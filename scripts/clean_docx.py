import docx
import re
import os

def clean_docx(filepath):
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return
        
    try:
        doc = docx.Document(filepath)
    except Exception as e:
        print(f"Failed to open {filepath}: {e}")
        return
        
    changed = False
    
    # We want to match github links or just clear text after "GitHub Link:"
    github_pattern = re.compile(r'https?://github\.com\S*')
    
    # Process paragraphs
    for p in doc.paragraphs:
        if 'github.com' in p.text.lower():
            for run in p.runs:
                if 'github.com' in run.text.lower():
                    run.text = github_pattern.sub('', run.text)
                    changed = True
            
            # Re-check paragraph text in case runs didn't split perfectly
            if github_pattern.search(p.text):
                p.text = github_pattern.sub('', p.text)
                changed = True
                
        # If it says "GitHub Link: [Link to GitHub Repository]" or similar
        if 'github link:' in p.text.lower():
            p.text = "GitHub Link: "
            changed = True
                
    # Process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if 'github.com' in p.text.lower():
                        for run in p.runs:
                            if 'github.com' in run.text.lower():
                                run.text = github_pattern.sub('', run.text)
                                changed = True
                        if github_pattern.search(p.text):
                            p.text = github_pattern.sub('', p.text)
                            changed = True
                    if 'github link:' in p.text.lower():
                        p.text = "GitHub Link: "
                        changed = True

    if changed:
        try:
            doc.save(filepath)
            print(f"Updated {filepath}")
        except Exception as e:
            print(f"Failed to save {filepath}: {e}")
    else:
        print(f"No changes in {filepath}")

clean_docx('docs/CAPSTONE_SUBMISSION_REPORT.docx')
clean_docx('docs/Capstone_Submission_Knowledge_Assistant.docx')
