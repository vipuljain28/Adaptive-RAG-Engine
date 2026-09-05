import docx
import re

def clean_headers_footers(filepath):
    try:
        doc = docx.Document(filepath)
        changed = False
        github_pattern = re.compile(r'https?://github\.com\S*')
        
        for section in doc.sections:
            # Check headers
            for p in section.header.paragraphs:
                if 'github' in p.text.lower():
                    print(f"Found github in header: {p.text}")
                    for run in p.runs:
                        if 'github' in run.text.lower():
                            run.text = ""
                            changed = True
                    p.text = p.text.replace('GitHub Link:', 'GitHub Link: ')
                    changed = True
            
            for table in section.header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if 'github' in p.text.lower():
                                print(f"Found github in header table: {p.text}")
                                p.text = p.text.replace(p.text, "")
                                changed = True
            
            # Check footers
            for p in section.footer.paragraphs:
                if 'github' in p.text.lower():
                    print(f"Found github in footer: {p.text}")
                    for run in p.runs:
                        if 'github' in run.text.lower():
                            run.text = ""
                            changed = True
                    changed = True
                    
        if changed:
            doc.save(filepath)
            print(f"Saved changes to headers/footers in {filepath}")
        else:
            print(f"No github text in headers/footers of {filepath}")
    except Exception as e:
        print(f"Error: {e}")

clean_headers_footers('docs/CAPSTONE_SUBMISSION_REPORT.docx')
