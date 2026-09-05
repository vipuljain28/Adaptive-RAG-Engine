import docx
from docx.shared import Inches
import os
import glob

def append_screenshots(docx_path, docs_folder, out_file):
    try:
        print(f"Opening {docx_path}...")
        doc = docx.Document(docx_path)
        
        doc.add_page_break()
        p = doc.add_paragraph()
        run = p.add_run('Appendix: Project Screenshots')
        run.bold = True
        run.font.size = docx.shared.Pt(18)
        
        png_files = glob.glob(os.path.join(docs_folder, '*.png'))
        if not png_files:
            print("No PNG files found in docs folder.")
            return

        for img_path in sorted(png_files):
            filename = os.path.basename(img_path)
            
            p2 = doc.add_paragraph()
            run2 = p2.add_run(filename.replace('.png', '').replace('_', ' '))
            run2.bold = True
            
            try:
                doc.add_picture(img_path, width=Inches(6.0))
                print(f"Added {filename}")
            except Exception as e:
                print(f"Could not add {filename}: {e}")
                
            doc.add_paragraph()

        doc.save(out_file)
        print(f"Successfully appended {len(png_files)} screenshots and saved {out_file}")
        
    except Exception as e:
        print(f"Error: {e}")

if __name__ == '__main__':
    docx_file = 'docs/CAPSTONE_SUBMISSION_REPORT.docx'
    docs_dir = 'docs'
    out_file = 'docs/FINAL_CAPSTONE_SUBMISSION_REPORT_WITH_IMAGES.docx'
    append_screenshots(docx_file, docs_dir, out_file)
